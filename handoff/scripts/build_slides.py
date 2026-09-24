# Fills the working deck (work.pptx, slides already duplicated and ordered) and writes out.pptx.
import copy
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY, GREY, CARD, INK = RGBColor(0x1B, 0x3A, 0x57), RGBColor(0x5B, 0x6B, 0x7A), RGBColor(0xEE, 0xF2, 0xF6), RGBColor(0x1E, 0x2A, 0x36)
OPTIONAL = "DIAPOSITIVE FACULTATIVE. Pour la retirer : clic droit sur sa miniature, puis « Supprimer la diapositive ». Aucune autre diapositive n'y renvoie.\n\n"

p = pptx.Presentation('work.pptx')
S = list(p.slides)


def named(slide):
    return {sh.name: sh for sh in slide.shapes}


def st(shape, text):
    """Replace a shape's text, keeping the formatting of its first run."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.text = text
    for para in tf.paragraphs[1:]:
        para._p.getparent().remove(para._p)


NOTES_TEMPLATE = S[23].notes_slide  # an original slide whose notes page has a text body


def notes(slide, text):
    ns = slide.notes_slide
    if ns.notes_text_frame is None:
        # This deck's notes master has no placeholders, so a new notes page is empty: copy an original one's shapes.
        tree = ns.shapes._spTree
        for el in NOTES_TEMPLATE.shapes._spTree.iterchildren():
            if el.tag.endswith('}sp'):
                tree.append(copy.deepcopy(el))
    ns.notes_text_frame.text = text


def fill(slide, pairs):
    n = named(slide)
    for name, text in pairs.items():
        st(n[name], text)


def cards(slide, title, subtitle, items):
    """Slides copied from 'Limites': six cards, heading + body."""
    heads = ['Text 3', 'Text 6', 'Text 9', 'Text 12', 'Text 15', 'Text 18']
    bodies = ['Text 4', 'Text 7', 'Text 10', 'Text 13', 'Text 16', 'Text 19']
    pairs = {'Text 0': title, 'Text 1': subtitle}
    for (h, b), hn, bn in zip(items, heads, bodies):
        pairs[hn] = h
        pairs[bn] = b
    fill(slide, pairs)


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, font='Calibri', align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = align
    r = para.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
    r.font.color.rgb = color
    return tb


# ---- 2 Plan -------------------------------------------------------------------------
fill(S[1], {
    'Text 11': 'Besoins et conception', 'Text 12': "les utilisateurs, puis l'architecture",
    'Text 15': 'Réalisation et démonstration', 'Text 16': 'le produit, ses versions, en images',
    'Text 19': 'Coûts et mise sur le marché', 'Text 20': "ce qu'il coûte, où l'héberger, pour qui",
})
notes(S[1], "Voici le plan. D'abord le problème et nos objectifs. Ensuite ce qui existe déjà. Puis les besoins et la conception. "
      "Nous montrerons le produit, ses versions et une démonstration. Nous parlerons ensuite de Sanad comme d'un produit : "
      "ce qu'il coûte, où l'héberger quand les données sont privées, et comment le mettre sur le marché. "
      "Et nous finirons par les résultats, les limites et la suite.")

# ---- 4 Objectifs --------------------------------------------------------------------
fill(S[3], {'Text 1': "Trois promesses, chacune avec un chiffre, fixées au lancement des sprints, avant de construire le produit."})
n4 = S[3].notes_slide.notes_text_frame.text.replace(
    "Nous avons écrit trois promesses chiffrées avant de coder.",
    "Nous avons écrit trois promesses chiffrées en juillet, au lancement des sprints, avant de construire le produit.")
notes(S[3], n4)

# ---- 6 Conduite du projet (optional) -----------------------------------------------
s = S[5]
n = named(s)
fill(s, {
    'Text 1': "Commencé en mars. Depuis juillet : un plan, des sprints, des tâches pour chacun, des versions.",
    'Text 8': 'Chacun sa partie', 'Text 9': 'Meriem : documents, mesure, écrans. Youssef : agent, mise en ligne.',
    'Text 13': "tests automatiques, puis relue par l'autre.",
})
pic = n['Image 0']
pic._element.getparent().remove(pic._element)
rows = [
    ('Mars 2026', "Étude du sujet : modèles de langage, RAG, outils existants."),
    ('Avril 2026', "Début du code, en local, sur nos machines."),
    ('20 juillet 2026', "Travail d'équipe structuré : dépôt GitHub commun, plan, sprints, tâches."),
    ('12 au 19 sept.', "Cinq versions publiées, chacune mesurée avant de sortir."),
]
for i, (when, what) in enumerate(rows):
    y = 1.72 + i * 0.86
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(y), Inches(8.2), Inches(0.72))
    box.adjustments[0] = 0.12
    box.fill.solid()
    box.fill.fore_color.rgb = CARD
    box.line.fill.background()
    box.shadow.inherit = False
    textbox(s, 0.95, y + 0.2, 2.1, 0.35, when, 16, True, NAVY, 'Cambria')
    textbox(s, 3.1, y + 0.22, 5.6, 0.35, what, 14, False, GREY)
notes(s, OPTIONAL + "Le projet a commencé en mars par l'étude du sujet, et le code en avril, sur nos machines. "
      "À partir du 20 juillet, nous avons travaillé comme une équipe professionnelle : un dépôt GitHub commun, un plan, "
      "des sprints, des tâches attribuées à chacun, et une livraison par versions. À droite, nos règles : une réunion chaque samedi, "
      "chacun sa partie, et aucune modification acceptée sans tests automatiques puis relecture par l'autre. "
      "En bas, le projet en chiffres : huit sprints, 146 modifications fusionnées, plus de 1 400 tests, cinq versions.")

# ---- 7 Besoins ----------------------------------------------------------------------
s = S[6]
n = named(s)
fill(s, {
    'Text 1': "Deux usages, pas deux rôles : préparer ses espaces, puis les interroger.",
    'Text 4': 'Interroger', 'Text 5': 'choisir un espace, poser sa question, ouvrir la source, voir les figures.',
    'Text 8': 'Préparer ses espaces', 'Text 9': "créer un espace, y déposer ses documents, lancer la synchronisation.",
    'Text 13': "un dossier de documents ; seul son propriétaire le modifie.",
    'Text 15': "Exemple : une chargée RH dépose le Code du travail dans son espace, puis demande la durée de la période d'essai d'un cadre.",
})
old = n['Image 0']
new = s.shapes.add_picture('usecase.png', old.left, old.top, height=old.height)
new.left = old.left + (old.width - new.width) // 2
old._element.addprevious(new._element)
old._element.getparent().remove(old._element)
notes(s, "Voici le diagramme des cas d'utilisation. Un seul acteur humain : l'utilisateur connecté. "
      "Il n'y a pas de rôle d'administrateur : chacun crée ses propres espaces et en est propriétaire, lui seul peut les modifier. "
      "Les espaces créés avant les comptes restent partagés, en lecture seule. Le modèle de langage est un acteur externe : "
      "il rédige les réponses et décrit les figures. Précision honnête : ces besoins viennent de notre expérience en entreprise, pas d'entretiens.\n\n"
      "Si le jury demande pourquoi il n'y a pas de rôles : la version 3.0 en avait trois (administrateur, curateur, lecteur). "
      "Nous les avons retirés en 3.1 pour une démonstration ouverte à l'inscription. Pour une entreprise, il faut les remettre (diapositive 19).")

# ---- 13 Versions (new) --------------------------------------------------------------
cards(S[12], 'Les versions', "Cinq versions en huit jours, chacune mesurée avant de sortir.", [
    ('v1.0.0 · 12 septembre', "espaces, synchronisation, réponses avec sources, refus, clarification, mémoire, rapports d'évaluation"),
    ('v1.0.1 · 12 septembre', "délais et reprises sur les appels au modèle, démarrage plus rapide, jeu de questions corrigé"),
    ('v2.0.0 · 13 septembre', "trace de la réponse, PowerPoint, PDF scannés, arabe de droite à gauche, avis, dossier surveillé"),
    ('v3.0.0 · 14 septembre', "comptes, dépôt depuis le navigateur, réponses en direct, tableau de bord, français, arabe, anglais"),
    ('v3.1.0 · 19 septembre', "espaces personnels, historique des conversations, espace créé depuis un dossier, limites de débit"),
    ('Suivante · 19 septembre', "figures : schémas et photos extraits, décrits et affichés avec la réponse"),
])
notes(S[12], "Depuis juillet, nous avons livré par versions. Chaque version ajoute un lot de fonctions, et chacune a passé la porte "
      "de publication avant de sortir : 36, 37, 38, 38 puis 39 réponses fondées sur 40, et 20 refus sur 20 à chaque fois. "
      "La 1.0 contient le cœur : espaces, synchronisation, réponses sourcées, refus. La 2 élargit les documents et ajoute l'arabe. "
      "La 3 rend Sanad utilisable à plusieurs : comptes, dépôt depuis le navigateur, réponses en direct. La 3.1 donne à chacun ses espaces "
      "et son historique. La dernière ajoute les figures.\n\nÀ savoir : la 3.0 avait trois rôles et une page d'administration ; "
      "la 3.1 les a retirés, et chaque personne connectée est désormais propriétaire de ses propres espaces.")

# ---- 15 Problèmes rencontrés (new, optional) ----------------------------------------
cards(S[14], 'Problèmes rencontrés', "Six vrais problèmes, et la solution technique de chacun.", [
    ('Un filigrane sur 201 pages', "Chaque page partait au modèle de mise en page : 30 min pour aucune figure. Solution : une image répétée est écartée."),
    ('Google retire notre modèle (27 août)', "Le premier vrai appel échoue. Solution : un nom de modèle fixé, vérifié par un vrai appel à chaque changement."),
    ("La description d'une image atteignait le vérificateur", "Une description inventée pouvait devenir une réponse. Solution : deux textes, la réponse ne lit que le document."),
    ('Le conteneur ne démarrait plus', "La bibliothèque torch pour carte graphique pèse plusieurs Go. Solution : version processeur imposée à la construction."),
    ("Le mot du refus s'affichait à l'écran", "Pendant la rédaction en direct. Solution : les 40 premiers caractères sont retenus un instant."),
    ('Une recherche dégradée, sans erreur', "Le préfixe du modèle E5 était oublié. Solution : un test vérifie le préfixe de chaque texte rangé."),
])
notes(S[14], OPTIONAL + "Six problèmes réels du projet. Un : le Code du travail porte le même filigrane sur ses 201 pages ; sans filtre, "
      "toute la loi partait au modèle de mise en page, environ trente minutes pour aucune figure. Deux : Google a retiré notre modèle le 27 août, "
      "et le premier vrai appel a échoué ; le nom est maintenant fixé et vérifié par un appel réel. Trois : la description générée d'une image "
      "arrivait au vérificateur ; une description inventée aurait pu transformer un refus honnête en réponse. Quatre : la version de torch pour "
      "carte graphique empêchait le conteneur de démarrer. Cinq : le mot de code du refus apparaissait une fraction de seconde pendant la rédaction "
      "en direct. Six : le modèle E5 sans son préfixe fonctionne, mais cherche mal, sans aucune erreur ; un test le vérifie.")

# ---- 16 Et si... (new, optional) ------------------------------------------------------
cards(S[15], "Et si nous ne l'avions pas fait ?", "Six choix discrets, et l'erreur que chacun évite.", [
    ('Sans parent_id', "On trouverait 500 caractères sans pouvoir lire l'article entier : réponses coupées, sources sans contexte."),
    ("Sans propriétaire sur l'espace", "Changer un numéro dans l'adresse ouvrirait l'espace d'un autre. Le serveur vérifie donc à chaque accès."),
    ('Sans effacer avant de réindexer', "Après une modification du fichier, Sanad citerait un texte qui n'y est plus."),
    ('Sans limite de reformulations', "L'agent pourrait chercher sans fin, et chaque essai coûte un appel au modèle."),
    ('Sans source obligatoire dans le code', "Si le modèle oubliait la source, une réponse sans preuve sortirait. Ici, elle ne peut pas être construite."),
    ('Sans questions gelées à l\'avance', "Nous aurions réglé Sanad sur l'examen : des scores flatteurs qui ne disent rien."),
])
notes(S[15], OPTIONAL + "Chaque case répond à la question : que se passerait-il sans ce choix ? Par exemple, parent_id est le lien entre un petit "
      "passage trouvé et sa section complète : sans lui, on trouve la bonne phrase mais on ne peut pas lire l'article autour. "
      "Et le lien entre un espace et son propriétaire : sans lui, n'importe qui pourrait lire l'espace d'un autre en changeant un numéro dans l'adresse.\n\n"
      "Problèmes qui pourraient encore arriver, si le jury demande : un document qui contient des consignes cachées pour le modèle ; "
      "des pics de mémoire jusqu'à 10 Go pendant l'extraction des figures ; une session qui reste valide jusqu'à son expiration après la suppression "
      "d'un compte ; des limites de débit gardées dans un seul processus, donc un seul serveur à la fois.")

# ---- 17 Exigences (new) -------------------------------------------------------------
cards(S[16], 'Exigences retenues', "Ce que Sanad doit faire, et les qualités qu'il doit avoir. Chacune est vérifiée.", [
    ('Fonctionnelle · répondre avec la source', "le fichier et la section sous chaque réponse : 39 réponses sur 39"),
    ("Fonctionnelle · refuser quand c'est absent", "20 questions hors documents sur 20 refusées, à chaque version"),
    ('Fonctionnelle · chacun ses espaces', "seul le propriétaire modifie un espace ; vérifié à chaque accès"),
    ('Non fonctionnelle · rapidité', "8,3 s par réponse pour 20 s visées ; 200 pages en moins de 10 min au repos"),
    ('Non fonctionnelle · confidentialité', "en mode local, rien ne sort ; en cloud, seuls la question et les passages"),
    ('Non fonctionnelle · coût', "logiciels libres, sans licence ; hébergement et modèle payés à l'usage"),
])
notes(S[16], "Les exigences que nous avons retenues, en deux familles. À gauche en haut, les fonctions : répondre avec la source, refuser "
      "quand la réponse est absente, et séparer les espaces de chacun. Les autres cases sont des qualités : la rapidité, la confidentialité "
      "et le coût. Chacune a une mesure ou un test ; les chiffres viennent des rapports d'évaluation.")

# ---- 18 Coûts (new, copied from Performances) ---------------------------------------
s = S[17]
n = named(s)
fill(s, {
    'Text 0': 'Ce que Sanad coûte', 'Text 1': "Des chiffres réels relevés sur Railway, et une estimation pour le modèle.",
    'Text 3': '7,60 $', 'Text 4': "hébergement réel, du 6 au 24 septembre, deux déploiements",
    'Text 6': '≈ 18 $ / mois', 'Text 7': "en ligne, au repos : surtout la mémoire (1,8 Go)",
    'Text 9': '≈ 1 centime', 'Text 10': "par question envoyée au modèle (estimation)",
})
chart = n['Chart 0']
rid = chart._element.xpath('.//c:chart/@r:id')
chart._element.getparent().remove(chart._element)
for r in rid:
    s.part.drop_rel(r)
rows = [
    ('Poste', 'Coût'),
    ('Logiciels : Python, Qdrant, Keycloak…', '0 $, tous libres'),
    ('GitHub et tests automatiques', '0 $, dépôt public'),
    ('Stockage : 2 Go sur le serveur', '≈ 0,30 $ / mois'),
    ('1 000 questions avec Gemini 3.6 Flash', '≈ 8 $'),
    ('1 000 questions avec Claude Haiku 4.5', '≈ 11 $'),
    ('Serveur GPU privé loué (L4, 24 Go)', '≈ 575 € / mois HT'),
]
tbl = s.shapes.add_table(len(rows), 2, Inches(5.3), Inches(1.85), Inches(7.33), Inches(0.5 * len(rows))).table
tbl.columns[0].width, tbl.columns[1].width = Inches(4.6), Inches(2.73)
for i, row in enumerate(rows):
    tbl.rows[i].height = Inches(0.5)
    for j, text in enumerate(row):
        c = tbl.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY if i == 0 else (CARD if i % 2 else RGBColor(0xFF, 0xFF, 0xFF))
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = c.margin_right = Inches(0.15)
        para = c.text_frame.paragraphs[0]
        para.text = ''
        r = para.add_run()
        r.text = text
        r.font.size, r.font.name, r.font.bold = Pt(14), 'Calibri', i == 0 or j == 1
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK
textbox(s, 5.3, 5.5, 7.33, 0.8, "Estimation par question : environ 8 000 jetons lus et 500 écrits. Prix officiels relevés le 24 septembre 2026 "
        "(Railway, Google, Anthropic, Scaleway).", 11.5, False, GREY)
notes(s, "Ce que Sanad coûte, avec des chiffres réels. Railway facture à la minute ce qui est vraiment utilisé. Nos deux déploiements ont coûté "
      "7,60 dollars en tout, du 6 au 24 septembre. Au repos, les trois services occupent environ 1,8 Go de mémoire : l'application, Keycloak et sa base. "
      "Un mois complet coûte donc environ 18 dollars, presque tout en mémoire. Le stockage, 2 Go, coûte une trentaine de centimes par mois. "
      "Les logiciels sont libres et GitHub est gratuit pour un dépôt public.\n\nPour le modèle, c'est une estimation, pas une mesure : "
      "environ 8 000 jetons lus et 500 écrits par question, soit moins d'un centime avec Gemini 3.6 Flash et environ un centime avec Claude Haiku 4.5. "
      "Attention : le prix de Gemini 3.6 Flash double le 1er janvier 2027. Claude Sonnet 5, plus fort, coûterait environ 2 à 3 centimes par question.\n\n"
      "Question probable : combien a coûté le modèle pendant le projet ? Réponse honnête : nous n'avons pas relevé le nombre de jetons ; "
      "le chiffre exact est dans la console Google AI Studio.")

# ---- 19 Hébergement des données privées (new) ---------------------------------------
cards(S[18], "Données privées : où faire tourner le modèle ?", "La recherche reste toujours chez le client. Seule la rédaction peut changer de place.", [
    ('Option 1 · une API cloud payante', "Claude ou Gemini payant : rapide, bonne qualité, pas d'entraînement sur vos données. Mais elles sortent du Maroc."),
    ('Option 2 · un serveur GPU partagé', "Un modèle ouvert sur une seule machine pour toute l'équipe, environ 575 € par mois en location. Rien ne sort."),
    ('Pas un modèle sur chaque ordinateur', "Ollama sur chaque poste : lent sans carte graphique, et au moins 8 Go de mémoire pour un modèle 7B."),
    ("Jamais l'offre gratuite", "Chez Google, elle peut servir à améliorer ses produits, et ses conditions interdisent d'y mettre des données confidentielles."),
    ('La loi 09-08', "Des données personnelles à l'étranger : seulement dans les cas prévus, sinon avec l'autorisation de la CNDP."),
    ('Notre recommandation', "Documents publics : API payante. Données personnelles : serveur privé. Et des rôles par espace."),
])
notes(S[18], "Quand les données sont privées, où faire tourner le modèle ? D'abord, un point important : dans Sanad, la recherche est déjà locale. "
      "Le découpage, les vecteurs et l'index tournent sur la machine du client. Seule la rédaction de la réponse appelle un modèle.\n\n"
      "Option 1, une API cloud payante. Anthropic n'entraîne pas ses modèles sur les données de ses offres commerciales, par défaut. "
      "Google non plus, sur son offre payante. Mais les passages sortent du pays. Option 2, un serveur avec une carte graphique, partagé par toute "
      "l'équipe : un seul modèle ouvert, et rien ne sort. Chez Scaleway, un serveur L4 de 24 Go coûte 0,79 euro de l'heure hors taxes, environ "
      "575 euros par mois en continu, moins s'il ne tourne qu'aux heures de bureau. Ce que nous déconseillons : un modèle sur chaque ordinateur, "
      "trop lent sans carte graphique, et l'offre gratuite de Gemini.\n\nLa loi 09-08 : des données personnelles ne partent à l'étranger que dans les "
      "cas prévus par les articles 43 et 44, sinon avec l'autorisation de la CNDP. Notre recommandation est donc hybride. "
      "Et pour une entreprise, il faut remettre des rôles par espace : propriétaire, éditeur, lecteur, plus un administrateur.\n\n"
      "Précision : nous n'avons pas encore mesuré le mode local sur les 60 questions. C'est la première chose à faire avant de le vendre.")

# ---- 20 Sanad comme produit (new, copied from Perspectives) -------------------------
fill(S[19], {
    'Text 0': 'Sanad comme produit', 'Text 1': "Pour qui, combien, et en combien de temps.",
    'Text 3': 'Pour qui',
    'Text 6': "PME avec de longs documents : RH, qualité, maintenance",
    'Text 9': "Là où une réponse doit citer sa source",
    'Text 11': 'Rentabilité',
    'Text 14': "≈ 30 $ par mois en ligne, 1 000 questions comprises",
    'Text 17': "Rentable dès 3 heures gagnées par mois, à 10 $ l'heure",
    'Text 20': "Une hypothèse à mesurer pendant un pilote",
    'Text 22': 'Mise sur le marché',
    'Text 25': "Dans 1 mois : un pilote chez un client, avec des rôles",
    'Text 28': "Dans 3 mois : version entreprise, mode local mesuré",
})
textbox(S[19], 0.7, 5.75, 11.93, 0.6, "Modèle économique possible : le code reste ouvert ; le client paie l'installation, l'hébergement et le suivi.", 15, False, NAVY)
notes(S[19], "Sanad vu comme un produit. Pour qui : les petites et moyennes entreprises qui travaillent avec de longs documents, là où une réponse doit citer sa source.\n\n"
      "Rentabilité : en ligne, Sanad coûte environ 30 dollars par mois, hébergement et mille questions compris. Il devient rentable dès qu'il fait gagner "
      "à l'équipe plus de temps que ce montant ne représente : à 10 dollars de l'heure, trois heures par mois suffisent. C'est une hypothèse, "
      "pas une mesure : nous ne citons pas les chiffres célèbres sur le temps passé à chercher, car leur source est fragile. Un pilote le mesurera.\n\n"
      "Mise sur le marché : le produit existe et il est mesuré. Dans un mois, un pilote chez un client, avec des rôles rétablis et un serveur privé. "
      "Dans trois mois, une version entreprise, avec le mode local mesuré sur les 60 questions. Ce sont nos estimations.\n\n"
      "Questions probables. La concurrence ? NotebookLM est gratuit, mais les documents partent chez Google et son refus n'est pas mesuré. "
      "Le risque principal ? Le fournisseur du modèle change : c'est arrivé le 27 août, et la porte de publication l'a détecté. "
      "Le modèle économique ? Code ouvert, et un service payant : installation, hébergement, suivi.")

# ---- 25 Perspectives ----------------------------------------------------------------
fill(S[24], {'Text 20': 'Des rôles par espace : propriétaire, éditeur, lecteur'})
notes(S[24], S[24].notes_slide.notes_text_frame.text.replace(
    "et donner un gestionnaire à chaque espace partagé, parce qu'aujourd'hui il faut passer par le serveur pour retirer un fichier",
    "et remettre des rôles par espace, propriétaire, éditeur et lecteur, parce qu'aujourd'hui un espace partagé n'a pas de responsable et qu'il faut passer par le serveur pour en retirer un fichier"))

# ---- page numbers -------------------------------------------------------------------
for i, slide in enumerate(S):
    for sh in slide.shapes:
        if sh.has_text_frame and abs(sh.left - Inches(12.23)) < Inches(0.05) and abs(sh.top - Inches(6.98)) < Inches(0.05):
            st(sh, str(i + 1))

p.save('out.pptx')
print('saved out.pptx with', len(S), 'slides')
