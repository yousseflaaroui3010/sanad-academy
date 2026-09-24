import docx, pptx, sys
d = docx.Document('Page1_PFA.docx')
with open('report.txt','w',encoding='utf-8') as f:
    for i,p in enumerate(d.paragraphs):
        if p.text.strip(): f.write(f"[{i}|{p.style.name}] {p.text}\n")
    f.write(f"\n=== TABLES: {len(d.tables)}\n")
    for ti,t in enumerate(d.tables):
        f.write(f"--- table {ti}\n")
        for r in t.rows: f.write(" | ".join(c.text.replace('\n',' / ') for c in r.cells)+"\n")
    f.write(f"\n=== sections {len(d.sections)}, inline images {len(d.inline_shapes)}\n")
p = pptx.Presentation('presentationkharia.pptx')
with open('slides.txt','w',encoding='utf-8') as f:
    f.write(f"size {p.slide_width} x {p.slide_height}\n")
    for i,s in enumerate(p.slides,1):
        f.write(f"\n##### SLIDE {i} (layout: {s.slide_layout.name})\n")
        for sh in s.shapes:
            t = sh.text_frame.text.replace('\n',' / ') if sh.has_text_frame else ''
            kind = sh.shape_type
            if sh.has_table if hasattr(sh,'has_table') else False:
                t = ' || '.join(' | '.join(c.text for c in r.cells) for r in sh.table.rows)
            f.write(f"  - [{kind}] {sh.name}: {t}\n")
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            f.write(f"  NOTES: {s.notes_slide.notes_text_frame.text[:600]}\n")
